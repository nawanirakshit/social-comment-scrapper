"""Export the last run's dataset to an Excel workbook.

Usage (from the project root):
    .venv/bin/python scripts/export_excel.py [output.xlsx] [dataset dir]

Defaults: results.xlsx in the project root, reading storage/datasets/default.
Sheets: Posts, Comments, and Errors (only when the run produced error rows).
Timestamps are converted to real Excel dates in UTC so they sort and filter.
"""

from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
EXCEL_CELL_LIMIT = 32_767

# (header, row key, column width, wrap text)
POST_COLUMNS = [
    ("Platform", "platform", 11, False),
    ("Author", "author", 22, False),
    ("Posted (UTC)", "timestamp", 18, False),
    ("Likes", "likes", 9, False),
    ("Comments", "commentsCount", 11, False),
    ("Shares", "shares", 9, False),
    ("Caption", "text", 70, True),
    ("Post URL", "url", 42, False),
    ("Image URL", "imageUrl", 42, False),
]
COMMENT_COLUMNS = [
    ("Author", "author", 22, False),
    ("Comment", "text", 80, True),
    ("Likes", "likes", 8, False),
    ("Replies", "replies", 9, False),
    ("Posted (UTC)", "timestamp", 18, False),
    ("Reply to", "parentId", 20, False),
    ("Comment ID", "id", 20, False),
    ("Post URL", "postUrl", 42, False),
]
ERROR_COLUMNS = [
    ("Platform", "platform", 11, False),
    ("Error", "error", 18, False),
    ("URL", "url", 60, False),
]

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="44546A")


def load_rows(dataset_dir: Path) -> list[dict]:
    files = sorted(f for f in dataset_dir.glob("*.json") if not f.name.startswith("__"))
    return [json.loads(f.read_text()) for f in files]


def to_cell(key: str, value):
    if value is None:
        return None
    if key == "timestamp" and isinstance(value, str):
        # Parser output is "2026-09-28T06:54:59Z", or "2026-09-27" from meta tags.
        try:
            return datetime.fromisoformat(value.replace("Z", "+00:00")).replace(tzinfo=None)
        except ValueError:
            return value
    if isinstance(value, str):
        # Excel rejects control characters and caps cells at 32,767 characters.
        return ILLEGAL_CHARACTERS_RE.sub("", value)[:EXCEL_CELL_LIMIT]
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)[:EXCEL_CELL_LIMIT]
    return value


def add_sheet(wb: Workbook, title: str, columns: list[tuple], rows: list[dict]) -> None:
    ws = wb.create_sheet(title)
    ws.append([header for header, *_ in columns])
    for row in rows:
        ws.append([to_cell(key, row.get(key)) for _, key, *_ in columns])

    for i, (_, key, width, wrap) in enumerate(columns, start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = width
        for cell in ws[letter][1:]:
            if key == "timestamp" and isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm"
            if wrap:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                cell.alignment = Alignment(vertical="top")
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def export(rows: list[dict], out: Path) -> dict[str, int]:
    by_type = {t: [r for r in rows if r.get("type") == t] for t in ("post", "comment", "error")}
    wb = Workbook()
    wb.remove(wb.active)
    add_sheet(wb, "Posts", POST_COLUMNS, by_type["post"])
    add_sheet(wb, "Comments", COMMENT_COLUMNS, by_type["comment"])
    if by_type["error"]:
        add_sheet(wb, "Errors", ERROR_COLUMNS, by_type["error"])
    wb.save(out)
    return {t: len(r) for t, r in by_type.items()}


def main() -> None:
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "results.xlsx"
    dataset_dir = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "storage" / "datasets" / "default"
    rows = load_rows(dataset_dir)
    if not rows:
        sys.exit(f"No dataset rows found in {dataset_dir}. Run the actor first.")
    try:
        counts = export(rows, out)
    except PermissionError:
        sys.exit(f"Can't write {out}: close it in Excel and run again.")
    print(f"Wrote {out}: {counts['post']} posts, {counts['comment']} comments, {counts['error']} errors.")


if __name__ == "__main__":
    main()
