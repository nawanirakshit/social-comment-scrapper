from datetime import datetime

from openpyxl import load_workbook

from scripts.export_excel import export


def test_export_sheets_and_cell_values(tmp_path):
    rows = [
        {"platform": "instagram", "type": "post", "url": "https://x/p/1", "author": "a",
         "text": "cap\x07tion", "likes": 5, "commentsCount": 2, "timestamp": "2026-09-27"},
        {"platform": "instagram", "type": "comment", "postUrl": "https://x/p/1", "id": "1",
         "author": "b", "text": "hi 👋", "likes": 0, "replies": None, "parentId": None,
         "timestamp": "2026-10-03T06:10:04Z"},
        {"platform": "facebook", "type": "error", "url": "https://fb/1", "error": "login_wall"},
    ]
    out = tmp_path / "r.xlsx"
    assert export(rows, out) == {"post": 1, "comment": 1, "error": 1}

    wb = load_workbook(out)
    assert wb.sheetnames == ["Posts", "Comments", "Errors"]
    post = {h.value: c.value for h, c in zip(wb["Posts"][1], wb["Posts"][2])}
    assert post["Caption"] == "caption"  # control character stripped
    assert post["Posted (UTC)"] == datetime(2026, 9, 27)
    comment = {h.value: c.value for h, c in zip(wb["Comments"][1], wb["Comments"][2])}
    assert comment["Comment"] == "hi 👋"
    assert comment["Posted (UTC)"] == datetime(2026, 10, 3, 6, 10, 4)


def test_no_errors_sheet_without_errors(tmp_path):
    out = tmp_path / "r.xlsx"
    export([{"type": "post", "author": "a"}], out)
    assert load_workbook(out).sheetnames == ["Posts", "Comments"]
